"""Read transaction rows from a converted M-Pesa statement workbook."""

from datetime import date, datetime
import re

from openpyxl import load_workbook


class StatementReader:
    """Capture completed transactions and statement totals from Excel."""

    TRANSACTION_COLUMNS = (
        "Receipt No.", "Completion Time", "Details", "Transaction Status",
        "Paid In", "Withdrawn",
    )

    def __init__(self, file_path) -> None:
        self.file_path = file_path
        self.stated_totals = {"Paid In": None, "Withdrawn": None}
        self.calculated_totals = {"Paid In": 0.0, "Withdrawn": 0.0}

    def get_transactions(self) -> list[dict]:
        transactions = []
        self.stated_totals = {"Paid In": None, "Withdrawn": None}
        self.calculated_totals = {"Paid In": 0.0, "Withdrawn": 0.0}
        workbook = load_workbook(self.file_path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                column_indexes = None
                for row in worksheet.iter_rows(values_only=True):
                    if self._is_totals_row(row):
                        self.stated_totals = {
                            "Paid In": self._to_monetary_value(row[1]),
                            "Withdrawn": self._to_monetary_value(row[2]),
                        }
                        continue
                    header_indexes = self._transaction_column_indexes(row)
                    if header_indexes is not None:
                        column_indexes = header_indexes
                        continue
                    if (column_indexes is not None
                            and self._is_transaction_row(row, column_indexes)
                            and self._has_complete_status(row, column_indexes)):
                        transaction = self._transaction_from_row(
                            row, column_indexes)
                        transactions.append(transaction)
                        self._add_to_calculated_totals(transaction)
        finally:
            workbook.close()
        combined_transactions = self._combine_transaction_charges(
            transactions
        )
        return list(reversed(combined_transactions))

    @staticmethod
    def completion_date(transaction: dict) -> date:
        value = transaction.get("Completion Time")
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            for date_format in (
                "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M:%S.%f",
                "%Y-%m-%d",
            ):
                try:
                    return datetime.strptime(
                        value.strip(), date_format).date()
                except ValueError:
                    continue
        raise ValueError(f"Invalid transaction completion date: {value!r}")

    @staticmethod
    def _combine_transaction_charges(transactions: list[dict]) -> list[dict]:
        by_receipt = {}
        for transaction in transactions:
            by_receipt.setdefault(transaction["Receipt No."], []).append(
                transaction)

        combined_transactions = []
        processed_receipts = set()
        for transaction in transactions:
            receipt = transaction["Receipt No."]
            if receipt in processed_receipts:
                continue
            matching = by_receipt[receipt]
            processed_receipts.add(receipt)
            if len(matching) == 1:
                combined = transaction.copy()
                combined["Transaction charge"] = None
            else:
                payment = max(
                    matching,
                    key=StatementReader._transaction_amount_magnitude,
                )
                combined = payment.copy()
                combined["Transaction charge"] = sum(
                    StatementReader._transaction_amount(candidate)
                    for candidate in matching if candidate is not payment
                )
            combined_transactions.append(combined)
        return combined_transactions

    @staticmethod
    def _transaction_amount(transaction: dict) -> float:
        for column in ("Paid In", "Withdrawn"):
            value = transaction.get(column)
            if value is not None:
                return value
        return 0.0

    @staticmethod
    def _transaction_amount_magnitude(transaction: dict) -> float:
        return abs(StatementReader._transaction_amount(transaction))

    def _add_to_calculated_totals(self, transaction: dict) -> None:
        for title in self.calculated_totals:
            value = transaction.get(title)
            if value is not None:
                self.calculated_totals[title] = round(
                    self.calculated_totals[title] + abs(value), 2)

    @classmethod
    def _transaction_column_indexes(cls, row) -> dict | None:
        indexes = {
            value.strip(): index
            for index, value in enumerate(row)
            if isinstance(value, str)
            and value.strip() in cls.TRANSACTION_COLUMNS
        }
        return indexes if {"Receipt No.", "Completion Time"} <= indexes.keys() else None

    @staticmethod
    def _is_transaction_row(row, indexes: dict) -> bool:
        receipt_index = indexes["Receipt No."]
        completion_index = indexes["Completion Time"]
        if len(row) <= max(receipt_index, completion_index):
            return False
        reference = row[receipt_index]
        if not isinstance(reference, str) or len(reference) != 10:
            return False
        if any(character.isspace() for character in reference):
            return False
        if any(character.isalpha() and not character.isupper()
               for character in reference):
            return False
        try:
            StatementReader.completion_date(
                {"Completion Time": row[completion_index]})
        except ValueError:
            return False
        return True

    @staticmethod
    def _has_complete_status(row, indexes: dict) -> bool:
        status_index = indexes.get("Transaction Status")
        if status_index is None or len(row) <= status_index:
            return False
        status = row[status_index]
        return (isinstance(status, str)
                and status.strip().casefold() in {"complete", "completed"})

    @classmethod
    def _transaction_from_row(cls, row, indexes: dict) -> dict:
        transaction = {
            title: row[index] if row[index] not in (None, "") else None
            for title, index in indexes.items()
        }
        is_paid_in = False
        for title in ("Paid In", "Withdrawn"):
            value = transaction.get(title)
            if value is not None:
                transaction[title] = cls._to_monetary_value(value)
                is_paid_in = is_paid_in or title == "Paid In"
        cleaned_details = cls._clean_details(
            transaction.get("Details"), is_paid_in
        )
        transaction["Details"] = cleaned_details
        transaction["Payee"] = cleaned_details
        return transaction

    @staticmethod
    def _is_totals_row(row) -> bool:
        return (len(row) >= 3 and isinstance(row[0], str)
                and row[0].strip().upper() == "TOTAL:")

    @staticmethod
    def _to_monetary_value(value) -> float | None:
        if value in (None, ""):
            return None
        return round(float(str(value).replace(",", "")), 2)

    @staticmethod
    def _clean_details(details, is_paid_in: bool) -> str:
        if details is None:
            return ""
        details = str(details).strip()
        if "***" in details:
            details = re.sub(r".*\d+\*+\d+\s+", "", details)
        if "-" in details:
            details = details.split("-", 1)[1].strip()
        if "Charge" not in details and is_paid_in:
            details = ""

        details = details.replace("_x000D_", " ")
        return re.sub(r"\s+", " ", details).strip()
