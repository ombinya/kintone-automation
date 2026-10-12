"""Read transaction records from an M-Pesa record workbook."""

from openpyxl import load_workbook


class MpesaRecordReader:
    """Capture transaction rows from the M-PESA1 or M-PESA2 worksheet."""

    WORKSHEET_NAMES = ("M-PESA2", "M-PESA1")
    WORKSHEET_KEY = "__worksheet_name__"
    ROW_NUMBER_KEY = "__row_number__"
    COLUMNS = (
        "Date",
        "Transaction Code",
        "Kintone Record ID",
        "Account Code",
        "Job Code",
        "JC3",
        "Payee/Payer",
        "Description",
        "Amount",
        "Trans Chr.",
        "Total Amount",
        "Balance",
        "Remarks",
    )

    def __init__(self, file_path) -> None:
        self.file_path = file_path
        self.worksheet_name: str | None = None

    def get_records(self) -> list[dict]:
        workbook = load_workbook(
            self.file_path,
            read_only=True,
            data_only=True,
        )
        try:
            worksheet = next(
                (
                    workbook[name]
                    for name in self.WORKSHEET_NAMES
                    if name in workbook.sheetnames
                ),
                None,
            )
            if worksheet is None:
                supported_names = ", ".join(self.WORKSHEET_NAMES)
                raise ValueError(
                    "The workbook must contain an M-Pesa worksheet named "
                    f"{supported_names}."
                )
            self.worksheet_name = worksheet.title

            header_found = False
            records = []
            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_col=1,
                    max_col=len(self.COLUMNS),
                    values_only=True,
                ),
                start=1,
            ):
                if not header_found:
                    header_found = tuple(row) == self.COLUMNS
                    continue

                if row[0] is None:
                    continue

                record = dict(zip(self.COLUMNS, row))
                if record.get("Total Amount") is None:
                    record["Total Amount"] = self._calculated_total(record)
                record[self.WORKSHEET_KEY] = worksheet.title
                record[self.ROW_NUMBER_KEY] = row_number
                records.append(record)

            if not header_found:
                raise ValueError(
                    "The M-Pesa worksheet does not contain the expected "
                    "columns. The first columns must be Date, Transaction "
                    "Code, Kintone Record ID, and Account Code, followed by "
                    "the remaining M-Pesa record columns."
                )
            return records
        finally:
            workbook.close()

    @staticmethod
    def _calculated_total(record):
        """Calculate the value of the spreadsheet's Amount + Charge formula."""
        amount = MpesaRecordReader._number(record.get("Amount"))
        charge = MpesaRecordReader._number(record.get("Trans Chr."))
        if amount is None:
            return None
        return amount + (charge or 0.0)

    @staticmethod
    def _number(value):
        if value in (None, ""):
            return None
        try:
            return float(str(value).replace(",", ""))
        except (TypeError, ValueError):
            return None
