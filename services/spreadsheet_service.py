"""Validation for M-Pesa record spreadsheets."""

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
import re
from uuid import uuid4

from openpyxl import load_workbook
from readers.mpesa_record_reader import MpesaRecordReader
from readers.statement_reader import StatementReader


@dataclass(frozen=True)
class SpreadsheetResult:
    file_path: Path
    records: tuple[dict, ...]
    worksheet_name: str
    month: int
    year: int


class SpreadsheetServiceError(ValueError):
    """An expected spreadsheet selection or validation failure."""


class SpreadsheetService:
    """Validate the selected spreadsheet against the accounting period."""

    EDITABLE_COLUMNS = {
        1: "Date",
        2: "Transaction Code",
        3: "Kintone Record ID",
        4: "Account Code",
        5: "Job Code",
        6: "JC3",
        7: "Payee/Payer",
        8: "Description",
        9: "Amount",
        10: "Trans Chr.",
        13: "Remarks",
    }
    FORMULA_COLUMNS = (11, 12, 14, 15, 16, 17, 18, 19, 20, 22, 23, 24, 25)

    def __init__(self, output_directory: Path) -> None:
        self.output_directory = Path(output_directory)

    def validate_file(
        self,
        file_path: str,
        month: int,
        year: int,
    ) -> SpreadsheetResult:
        spreadsheet = Path(file_path)
        if not spreadsheet.is_file():
            raise SpreadsheetServiceError(
                "Select an M-Pesa records spreadsheet first."
            )
        if spreadsheet.suffix.casefold() != ".xlsx":
            raise SpreadsheetServiceError(
                "The M-Pesa records file must be an .xlsx spreadsheet."
            )

        period_token = f"{year}{month:02d}"
        if re.search(rf"(?<!\d){period_token}(?!\d)", spreadsheet.stem) is None:
            raise SpreadsheetServiceError(
                f"The filename must contain {period_token} for the selected "
                f"month and year. For example: "
                f"'{period_token}MPESA 1 Record.xlsx' or "
                f"'MPESA 1 Record {period_token}.xlsx'."
            )

        try:
            reader = MpesaRecordReader(spreadsheet)
            records = reader.get_records()
        except ValueError as error:
            raise SpreadsheetServiceError(str(error)) from error
        except Exception as error:
            raise SpreadsheetServiceError(
                "The spreadsheet could not be read. Confirm that it is a "
                "valid, closed .xlsx workbook and try again."
            ) from error

        return SpreadsheetResult(
            file_path=spreadsheet,
            records=tuple(records),
            worksheet_name=reader.worksheet_name or "",
            month=month,
            year=year,
        )

    def save_updated_copy(
        self,
        source_file: str,
        records,
        assignments,
        missing_statement_transactions=(),
    ) -> Path:
        source = Path(source_file)
        if not source.is_file():
            raise SpreadsheetServiceError(
                "The selected spreadsheet no longer exists."
            )
        if not self.output_directory.is_dir():
            raise SpreadsheetServiceError(
                f"Output directory not found: {self.output_directory}"
            )

        output = self.output_directory / f"{source.stem}_updated.xlsx"
        temporary_output = output.with_name(
            f".{output.stem}.{uuid4().hex}.xlsx"
        )
        workbook = None
        try:
            workbook = load_workbook(source, data_only=False)
            self._rewrite_editable_cells(
                workbook,
                records,
                assignments,
                missing_statement_transactions,
            )

            calculation = getattr(workbook, "calculation", None)
            if calculation is not None:
                calculation.calcMode = "auto"
                calculation.fullCalcOnLoad = True
                calculation.forceFullCalc = True

            workbook.save(temporary_output)
            workbook.close()
            workbook = None
            temporary_output.replace(output)
            return output
        except SpreadsheetServiceError:
            raise
        except PermissionError as error:
            raise SpreadsheetServiceError(
                "The updated spreadsheet could not be saved. Close the file "
                "in Excel and try again."
            ) from error
        except Exception as error:
            raise SpreadsheetServiceError(
                "The updated spreadsheet could not be created."
            ) from error
        finally:
            if workbook is not None:
                workbook.close()
            if temporary_output.exists():
                temporary_output.unlink()

    def save_kintone_ids_copy(
        self,
        source_file: str,
        records,
        kintone_id_assignments,
        allow_replacement: bool = False,
    ) -> Path:
        source = Path(source_file)
        if not source.is_file():
            raise SpreadsheetServiceError(
                "The selected spreadsheet no longer exists."
            )
        if not self.output_directory.is_dir():
            raise SpreadsheetServiceError(
                f"Output directory not found: {self.output_directory}"
            )

        if source.stem.endswith("_kintone_updated"):
            output = source
        else:
            output = self.output_directory / (
                f"{source.stem}_kintone_updated.xlsx"
            )
        temporary_output = output.with_name(
            f".{output.stem}.{uuid4().hex}.xlsx"
        )
        workbook = None
        try:
            workbook = load_workbook(source, data_only=False)
            self._rewrite_editable_cells(
                workbook,
                records,
                assignments=(),
                kintone_id_assignments=kintone_id_assignments,
                allow_kintone_id_replacement=allow_replacement,
            )

            calculation = getattr(workbook, "calculation", None)
            if calculation is not None:
                calculation.calcMode = "auto"
                calculation.fullCalcOnLoad = True
                calculation.forceFullCalc = True

            workbook.save(temporary_output)
            workbook.close()
            workbook = None
            temporary_output.replace(output)
            return output
        except SpreadsheetServiceError:
            raise
        except PermissionError as error:
            raise SpreadsheetServiceError(
                "The Kintone-updated spreadsheet could not be saved. Close "
                "the file in Excel and try again."
            ) from error
        except Exception as error:
            raise SpreadsheetServiceError(
                "The Kintone-updated spreadsheet could not be created."
            ) from error
        finally:
            if workbook is not None:
                workbook.close()
            if temporary_output.exists():
                temporary_output.unlink()

    @classmethod
    def _rewrite_editable_cells(
        cls,
        workbook,
        records,
        assignments,
        missing_statement_transactions=(),
        kintone_id_assignments=(),
        allow_kintone_id_replacement: bool = False,
    ) -> None:
        worksheet, header_row = cls._record_worksheet(workbook)
        first_data_row = header_row + 1
        total_row = cls._total_row(worksheet, first_data_row)
        final_prepared_row = total_row - 1

        cls._validate_formula_rows(
            worksheet,
            first_data_row,
            final_prepared_row,
        )
        cls._verify_source_records(
            worksheet,
            records,
            first_data_row,
            final_prepared_row,
        )

        required_row_count = len(records) + len(
            missing_statement_transactions
        )
        prepared_row_count = final_prepared_row - first_data_row + 1
        if required_row_count > prepared_row_count:
            raise SpreadsheetServiceError(
                "The spreadsheet has "
                f"{prepared_row_count} formula-prepared data rows, but "
                f"{required_row_count} records need to be written. Extend the "
                "template formulas and validate the spreadsheet again."
            )

        updated_records = [dict(record) for record in records]
        updated_by_source = {
            cls._record_source_key(record): record
            for record in updated_records
        }
        for assignment in assignments:
            source_key = cls._record_source_key(
                assignment.spreadsheet_record
            )
            record = updated_by_source.get(source_key)
            if record is None:
                raise SpreadsheetServiceError(
                    "A transaction-code assignment could not be mapped back "
                    "to its validated spreadsheet row. Validate the file "
                    "again and retry."
                )

            receipt_number = assignment.statement_transaction.get(
                "Receipt No."
            )
            existing_code = record.get("Transaction Code")
            if existing_code not in (None, "") and not cls._same_code(
                existing_code,
                receipt_number,
            ):
                raise SpreadsheetServiceError(
                    "A transaction-code cell changed after validation. "
                    "Validate the spreadsheet again before reconciling."
                )
            record["Transaction Code"] = receipt_number

        for source_record, kintone_record_id in kintone_id_assignments:
            source_key = cls._record_source_key(source_record)
            record = updated_by_source.get(source_key)
            if record is None:
                raise SpreadsheetServiceError(
                    "A Kintone Record ID could not be mapped back to its "
                    "validated spreadsheet row. Validate the spreadsheet "
                    "again and retry."
                )
            if kintone_record_id in (None, ""):
                raise SpreadsheetServiceError(
                    "A matched Kintone record does not have a Record ID."
                )

            existing_id = record.get("Kintone Record ID")
            if (
                not allow_kintone_id_replacement
                and existing_id not in (None, "")
                and str(existing_id).strip()
                != str(kintone_record_id).strip()
            ):
                raise SpreadsheetServiceError(
                    "A spreadsheet row already contains a different Kintone "
                    "Record ID. Review the record and validate the "
                    "spreadsheet again."
                )
            record["Kintone Record ID"] = str(kintone_record_id).strip()

        updated_records.extend(
            cls._record_from_statement(transaction)
            for transaction in missing_statement_transactions
        )

        ordered_records = [
            record
            for _, record in sorted(
                enumerate(updated_records),
                key=lambda item: (
                    cls._record_date(item[1]),
                    item[0],
                ),
            )
        ]

        for row_number in range(first_data_row, final_prepared_row + 1):
            for column in cls.EDITABLE_COLUMNS:
                worksheet.cell(row_number, column).value = None

        for row_number, record in enumerate(
            ordered_records,
            start=first_data_row,
        ):
            for column, field_name in cls.EDITABLE_COLUMNS.items():
                worksheet.cell(row_number, column).value = record.get(
                    field_name
                )

    @staticmethod
    def _record_worksheet(workbook):
        for worksheet_name in MpesaRecordReader.WORKSHEET_NAMES:
            if worksheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[worksheet_name]
            for row_number in range(1, worksheet.max_row + 1):
                headings = tuple(
                    worksheet.cell(row_number, column).value
                    for column in range(1, len(MpesaRecordReader.COLUMNS) + 1)
                )
                if headings == MpesaRecordReader.COLUMNS:
                    return worksheet, row_number
        raise SpreadsheetServiceError(
            "The M-Pesa worksheet header could not be found while preparing "
            "the updated spreadsheet."
        )

    @staticmethod
    def _total_row(worksheet, first_data_row: int) -> int:
        for row_number in range(first_data_row, worksheet.max_row + 1):
            value = worksheet.cell(row_number, 9).value
            if (
                isinstance(value, str)
                and value.upper().startswith("=SUBTOTAL(")
            ):
                return row_number
        raise SpreadsheetServiceError(
            "The spreadsheet total row could not be found. The existing "
            "formula-prepared row range cannot be determined safely."
        )

    @classmethod
    def _validate_formula_rows(
        cls,
        worksheet,
        first_data_row: int,
        final_prepared_row: int,
    ) -> None:
        for row_number in range(first_data_row, final_prepared_row + 1):
            missing_formulas = [
                worksheet.cell(row_number, column).coordinate
                for column in cls.FORMULA_COLUMNS
                if worksheet.cell(row_number, column).data_type != "f"
            ]
            if missing_formulas:
                raise SpreadsheetServiceError(
                    "The spreadsheet contains a prepared row with missing "
                    f"formulas: {', '.join(missing_formulas)}. Restore the "
                    "template formulas and validate the spreadsheet again."
                )

    @classmethod
    def _verify_source_records(
        cls,
        worksheet,
        records,
        first_data_row: int,
        final_prepared_row: int,
    ) -> None:
        source_rows = set()
        for record in records:
            worksheet_name, row_number = cls._record_source_key(record)
            if worksheet_name != worksheet.title or not (
                first_data_row <= row_number <= final_prepared_row
            ):
                raise SpreadsheetServiceError(
                    "A spreadsheet record is outside the formula-prepared "
                    "data rows. Validate the file again before reconciling."
                )
            source_rows.add(row_number)

            for column, field_name in cls.EDITABLE_COLUMNS.items():
                cell = worksheet.cell(row_number, column)
                if cell.data_type == "f":
                    raise SpreadsheetServiceError(
                        f"Formula cell {cell.coordinate} is in an editable "
                        "spreadsheet column and will not be overwritten."
                    )
                if not cls._same_cell_value(
                    cell.value,
                    record.get(field_name),
                ):
                    raise SpreadsheetServiceError(
                        f"Cell {cell.coordinate} changed after the spreadsheet "
                        "was validated. Validate it again before reconciling."
                    )

        for row_number in range(first_data_row, final_prepared_row + 1):
            if row_number in source_rows:
                continue
            populated_cells = [
                worksheet.cell(row_number, column).coordinate
                for column in cls.EDITABLE_COLUMNS
                if worksheet.cell(row_number, column).value not in (None, "")
            ]
            if populated_cells:
                raise SpreadsheetServiceError(
                    "A partially completed row was not captured as a record: "
                    f"{', '.join(populated_cells)}. Complete or clear that "
                    "row, then validate the spreadsheet again."
                )

    @staticmethod
    def _record_source_key(record) -> tuple[str, int]:
        worksheet_name = record.get(MpesaRecordReader.WORKSHEET_KEY)
        row_number = record.get(MpesaRecordReader.ROW_NUMBER_KEY)
        if not worksheet_name or row_number is None:
            raise SpreadsheetServiceError(
                "A spreadsheet record could not be mapped back to its source "
                "row. Validate the file again and retry."
            )
        try:
            return str(worksheet_name), int(row_number)
        except (TypeError, ValueError) as error:
            raise SpreadsheetServiceError(
                "A spreadsheet record has an invalid source row number."
            ) from error

    @staticmethod
    def _record_date(record):
        value = record.get("Date")
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            for date_format in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
                try:
                    return datetime.strptime(
                        value.strip(),
                        date_format,
                    ).date()
                except ValueError:
                    continue
        raise SpreadsheetServiceError(
            f"Invalid spreadsheet Date value: {value!r}."
        )

    @staticmethod
    def _record_from_statement(transaction) -> dict:
        statement_amount = next(
            (
                transaction.get(field_name)
                for field_name in ("Paid In", "Withdrawn")
                if transaction.get(field_name) is not None
            ),
            None,
        )
        return {
            "Date": datetime.combine(
                StatementReader.completion_date(transaction),
                time.min,
            ),
            "Transaction Code": transaction.get("Receipt No."),
            "Kintone Record ID": None,
            "Account Code": None,
            "Job Code": None,
            "JC3": None,
            "Payee/Payer": transaction.get("Payee") or None,
            "Description": None,
            "Amount": SpreadsheetService._opposite_number(statement_amount),
            "Trans Chr.": SpreadsheetService._opposite_number(
                transaction.get("Transaction charge")
            ),
            "Remarks": None,
        }

    @staticmethod
    def _opposite_number(value):
        if value is None:
            return None
        try:
            result = -float(str(value).replace(",", ""))
        except (TypeError, ValueError) as error:
            raise SpreadsheetServiceError(
                f"Invalid statement monetary value: {value!r}."
            ) from error
        return 0.0 if abs(result) < 0.005 else result

    @staticmethod
    def _same_code(left, right) -> bool:
        if left in (None, "") or right in (None, ""):
            return left in (None, "") and right in (None, "")
        return str(left).strip().casefold() == str(right).strip().casefold()

    @staticmethod
    def _same_cell_value(left, right) -> bool:
        if left in (None, "") and right in (None, ""):
            return True
        if isinstance(left, datetime) and isinstance(right, date):
            if isinstance(right, datetime):
                return left == right
            return left.date() == right
        if isinstance(right, datetime) and isinstance(left, date):
            return right.date() == left
        return left == right
